#!/usr/bin/env python3
"""Builds the .xlsx fixtures of test/structured-data.test.js. Run once; the results are committed, so the tests need no Python.

  basic.xlsx     written by openpyxl (a producer other than ours): text, numbers, long IDs, dates, a formula without a cached
                 value, a merged range, a comment, a defined name, a hidden sheet, a sparse far-away cell
  excelish.xlsx  written by hand the way Excel writes: shared strings, cached formula values, a shared formula group, an
                 error value, an inline string, a boolean, escaped control characters (_x0001_), date and time styles
  date1904.xlsx  the 1904 date system
"""
import datetime, os, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))


def basic():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws.title = 'Q1 Sales'
    ws.append(['Name', 'Qty', 'Price', 'ID', 'Shipped', 'Note'])
    ws.append(['Müller & Söhne', 3, 19.9, '007', datetime.datetime(2024, 3, 5, 14, 30, 0), 'zwei  Leerzeichen, "Anführungszeichen"'])
    ws.append(['Zoë 😀', 12, 0.1, '00123456789012345678', datetime.datetime(2024, 12, 31), 'Zeile1\nZeile2'])
    ws.append(['Total', '=SUM(B2:B3)', None, None, None, None])
    ws.append([])
    ws.append(['merged title'])
    ws.merge_cells('A6:C6')
    ws['A1'].comment = Comment('Column titles', 'tg-data')
    ws['G2'] = 12345678901234567890  # more digits than a double keeps
    ws['H2'] = True
    ws['I2'] = datetime.time(9, 5, 30)
    ws['Z100'] = 'far away'
    hidden = wb.create_sheet('Hidden')
    hidden['A1'] = 'secret'
    hidden.sheet_state = 'hidden'
    wb.defined_names['Prices'] = DefinedName('Prices', attr_text="'Q1 Sales'!$C$2:$C$3")
    wb.save(os.path.join(HERE, 'basic.xlsx'))


CT = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
RELS = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>'
WBRELS = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>'
# xf 0 general, 1 = builtin 14 (date), 2 = builtin 22 (date+time), 3 = builtin 21 (time), 4 = custom "0.00 ""m""" (not a date)
STYLES = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><numFmts count="1"><numFmt numFmtId="164" formatCode="0.00&quot; m&quot;"/></numFmts><fonts count="1"><font><sz val="11"/></font></fonts><fills count="1"><fill><patternFill patternType="none"/></fill></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0"/></cellStyleXfs><cellXfs count="5"><xf numFmtId="0"/><xf numFmtId="14" applyNumberFormat="1"/><xf numFmtId="22" applyNumberFormat="1"/><xf numFmtId="21" applyNumberFormat="1"/><xf numFmtId="164" applyNumberFormat="1"/></cellXfs></styleSheet>'
SST = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="4" uniqueCount="4"><si><t>Item</t></si><si><t xml:space="preserve">  padded  </t></si><si><r><t>Rich </t></r><r><rPr><b/></rPr><t>text</t></r></si><si><t>ctl_x0001_char _x005F_x0041_</t></si></si></sst>'.replace('</si></si>', '</si>')
SHEET = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c><c r="D1" t="s"><v>3</v></c></row>
<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>2.5</v></c><c r="C2"><f>A2+B2</f><v>3.5</v></c><c r="D2" t="b"><v>1</v></c><c r="E2" s="4"><v>1.5</v></c></row>
<row r="3"><c r="A3"><v>0.1</v></c><c r="B3"><v>1E-3</v></c><c r="C3"><f t="shared" ref="C3:C4" si="0">A3*2</f><v>0.2</v></c><c r="D3" t="e"><v>#DIV/0!</v></c></row>
<row r="4"><c r="A4"><v>100000000000000000000</v></c><c r="C4"><f t="shared" si="0"/><v>200000000000000000000</v></c><c r="D4" t="inlineStr"><is><t>inline</t></is></c><c r="E4" t="str"><f>"x"&amp;"y"</f><v>xy</v></c><c r="F4"><v>12345678901234567890</v></c></row>
<row r="5"><c r="A5" s="1"><v>45356</v></c><c r="B5" s="2"><v>45356.604166666664</v></c><c r="C5" s="3"><v>0.5</v></c><c r="D5" s="1"><v>59</v></c><c r="E5" s="1"><v>61</v></c><c r="F5" s="1"><v>60</v></c></row>
<row r="6" hidden="1"><c r="A6" t="s"><v>0</v></c></row>
</sheetData><mergeCells count="1"><mergeCell ref="A7:B7"/></mergeCells></worksheet>'''


def excelish(name, date1904):
    wbxml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">%s<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>' % ('<workbookPr date1904="1"/>' if date1904 else '')
    with zipfile.ZipFile(os.path.join(HERE, name), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', CT)
        z.writestr('_rels/.rels', RELS)
        z.writestr('xl/workbook.xml', wbxml)
        z.writestr('xl/_rels/workbook.xml.rels', WBRELS)
        z.writestr('xl/styles.xml', STYLES)
        z.writestr('xl/sharedStrings.xml', SST)
        z.writestr('xl/worksheets/sheet1.xml', SHEET)


if __name__ == '__main__':
    basic()
    excelish('excelish.xlsx', False)
    excelish('date1904.xlsx', True)
    print('written to', HERE)
